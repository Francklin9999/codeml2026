"""Deterministic extraction and evidence rendering for the NOVA corpus.

The module intentionally keeps extraction local and dependency-light.  PDF text is
read with ``pypdf`` and spreadsheets with ``openpyxl``; both failures are fatal so
that a source can never disappear silently from the evidence inventory.
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import shutil
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path, PurePosixPath
from typing import Any, Iterable

from openpyxl import load_workbook
from pypdf import PdfReader


CORPUS_PREFIX = "Projet360_NOVA_ETUDIANTS/"
SUPPORTED_SUFFIXES = {".eml", ".pdf", ".xlsx", ".txt", ".csv", ".md", ".png"}
MEETING_TIME = re.compile(r"^(?P<time>\d{2}:\d{2})\s+(?P<text>.+)$")
TEAMS_TIME = re.compile(r"^\[?(?P<time>\d{1,2}:\d{2})\]?\s*[-–:]?\s*(?P<text>.+)$")
TICKET_COMMENT = re.compile(
    r"^(?P<day>\d{1,2})\s+(?P<month>janv|f[ée]vr|mars|avr|mai|juin|juil|ao[uû]t|sept|oct|nov|d[ée]c)"
    r"(?:\.|embre|ier|il|let|ût)?\s+(?P<time>\d{1,2}:\d{2})\s*[-–:]\s*(?P<text>.+)$",
    re.IGNORECASE,
)
MONTHS = {
    "janv": 1, "févr": 2, "fevr": 2, "mars": 3, "avr": 4, "mai": 5,
    "juin": 6, "juil": 7, "août": 8, "aout": 8, "sept": 9, "oct": 10,
    "nov": 11, "déc": 12, "dec": 12,
}


@dataclass
class Unit:
    id: str
    file: str
    kind: str
    text: str
    anchor: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)

    def as_dict(self, evidence_file: str) -> dict[str, Any]:
        return {
            "id": self.id,
            "file": self.file,
            "kind": self.kind,
            "text": self.text,
            "anchor": self.anchor or anchor_for(self.id),
            "evidence_file": evidence_file,
            **self.metadata,
        }


def decode_text(raw: bytes) -> tuple[str, str]:
    """Decode a textual source with a documented, deterministic fallback."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            pass
    raise UnicodeDecodeError("utf-8/cp1252", raw, 0, len(raw), "unsupported source encoding")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def anchor_for(locator: str) -> str:
    normalized = unicodedata.normalize("NFKD", locator).encode("ascii", "ignore").decode("ascii")
    return "loc-" + re.sub(r"[^A-Za-z0-9_-]+", "-", normalized).strip("-").lower()


def evidence_name(relative_path: str) -> str:
    return relative_path.replace("/", "__") + ".html"


def _safe_members(archive: zipfile.ZipFile) -> list[zipfile.ZipInfo]:
    members: list[zipfile.ZipInfo] = []
    for info in archive.infolist():
        if info.is_dir():
            continue
        posix = PurePosixPath(info.filename)
        if posix.is_absolute() or ".." in posix.parts or not info.filename.startswith(CORPUS_PREFIX):
            raise ValueError(f"Unsafe or unexpected ZIP member: {info.filename}")
        relative = PurePosixPath(info.filename[len(CORPUS_PREFIX):])
        if relative.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError(f"Unsupported corpus format: {relative}")
        members.append(info)
    return sorted(members, key=lambda item: item.filename)


def _source_class(relative_path: str) -> tuple[str, str, str | None]:
    name = PurePosixPath(relative_path).name
    if relative_path == "08_Archives_et_documents_connexes/Courriel_archive_17sept.eml":
        return "exact-duplicate", "Copie exacte du courriel E12; ne constitue pas une confirmation indépendante.", "email-e12"
    if name == "INV-778_Projet_ORION.pdf":
        return "other-project", "Facture du projet ORION, hors périmètre NOVA.", None
    if name in {"Invitation_Formation_Excel.txt", "Newsletter_Boreal_Septembre.txt", "Notes_personnelles_quelquun.txt"}:
        return "noise", "Document connexe sans autorité opérationnelle pour NOVA.", None
    if name in {"README.txt", "MANIFEST.csv"}:
        return "noise", "Métadonnée administrative du jeu de données.", None
    if name.endswith(".png"):
        return "image-only-evidence", "Capture à lire avec sa transcription et la source datée associée.", None
    attachment_groups = {
        "Architecture_NOVA_v1.pdf": "attachment-architecture-v1",
        "Architecture_NOVA_v2.pdf": "attachment-architecture-v2",
        "INV-003.pdf": "attachment-inv-003",
        "CR-04_Optimisation_mobile_BROUILLON.pdf": "attachment-cr-04",
        "Rapport_Statut_21sept.pdf": "attachment-rapport-statut",
    }
    if name in attachment_groups:
        return "attachment-copy", "Copie octet pour octet d'une pièce jointe EML; ne constitue pas une confirmation indépendante.", attachment_groups[name]
    if name in {
        "Charte_Projet_NOVA_v1.txt", "Plan_Projet_NOVA_v2.xlsx", "Plan_Projet_NOVA_v3_12sept.xlsx",
        "Rapport_Statut_21sept.pdf", "Registre_Risques_29sept.xlsx", "Plan_NOVA_preliminaire_juin.xlsx",
    }:
        return "derived/stale", "Document dérivé ou daté; à confronter aux décisions et validations plus récentes.", None
    return "primary", "Source de travail; son autorité et sa date doivent être évaluées au niveau du fait.", None


def _information_date(relative_path: str) -> str | None:
    return {
        "04_Documents_projet/Registre_Risques_29sept.xlsx": "2026-09-09",
        "04_Documents_projet/Rapport_Statut_21sept.pdf": "2026-09-21",
        "04_Documents_projet/Plan_Projet_NOVA_v3_12sept.xlsx": "2026-09-12",
        "04_Documents_projet/Plan_Projet_NOVA_v2.xlsx": None,
        "08_Archives_et_documents_connexes/Plan_NOVA_preliminaire_juin.xlsx": "2026-06",
        "04_Documents_projet/Charte_Projet_NOVA_v1.txt": "2026-07-07",
    }.get(relative_path)


def _extract_email(raw: bytes, relative: str, attachments_dir: Path) -> tuple[str, list[Unit], dict[str, Any]]:
    message = BytesParser(policy=policy.default).parsebytes(raw)
    doc_id = Path(relative).stem.split("_")[0]
    header_lines = [f"{key}: {message.get(key, '')}" for key in ("From", "To", "Date", "Subject")]
    bodies: list[str] = []
    attachments: list[dict[str, Any]] = []
    for part in message.walk():
        filename = part.get_filename()
        if filename:
            payload = part.get_payload(decode=True) or b""
            safe_name = Path(filename).name
            output_name = f"{Path(relative).stem}__{safe_name}"
            (attachments_dir / output_name).write_bytes(payload)
            attachments.append({"name": safe_name, "saved_as": output_name, "bytes": len(payload), "sha256": sha256(payload)})
            continue
        if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
            try:
                bodies.append(part.get_content())
            except (LookupError, UnicodeDecodeError):
                payload = part.get_payload(decode=True) or b""
                bodies.append(decode_text(payload)[0])
    body = "\n\n".join(item.strip() for item in bodies if item.strip())
    units = [
        Unit(f"{doc_id}#headers", relative, "email-headers", "\n".join(header_lines)),
        Unit(f"{doc_id}#body", relative, "email-body", body),
    ]
    for index, paragraph in enumerate((p.strip() for p in re.split(r"\n\s*\n", body)), 1):
        if paragraph:
            units.append(Unit(f"{doc_id}#p{index}", relative, "email-paragraph", paragraph))
    if attachments:
        units.append(Unit(f"{doc_id}#attachments", relative, "email-attachments", "\n".join(a["name"] for a in attachments)))
    text = "\n".join(header_lines) + "\n\n" + body
    if attachments:
        text += "\n\nAttachments:\n" + "\n".join(f"- {a['name']} [{a['sha256']}]" for a in attachments)
    return text, units, {"attachments": attachments}


def _extract_pdf(raw: bytes, relative: str) -> tuple[str, list[Unit], dict[str, Any]]:
    reader = PdfReader(io.BytesIO(raw))
    pages: list[str] = []
    units: list[Unit] = []
    name = Path(relative).name
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if not text:
            raise ValueError(f"PDF page has no extractable text: {relative} page {number}")
        pages.append(f"=== page {number} ===\n{text}")
        units.append(Unit(f"{name}#p{number}", relative, "pdf-page", text, metadata={"page": number}))
    return "\n\n".join(pages), units, {"pages": len(pages)}


def _value_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    return "" if value is None else str(value)


def _extract_xlsx(raw: bytes, relative: str) -> tuple[str, list[Unit], dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    name = Path(relative).name
    multiple = len(workbook.sheetnames) > 1
    lines: list[str] = []
    units: list[Unit] = []
    comments = 0
    for sheet in workbook.worksheets:
        lines.append(f"=== sheet {sheet.title} ===")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and cell.comment is None:
                    continue
                value = _value_text(cell.value)
                locator = f"{name}!{sheet.title}!{cell.coordinate}" if multiple else f"{name}!{cell.coordinate}"
                line = f"{cell.coordinate}={value}"
                metadata: dict[str, Any] = {"sheet": sheet.title, "cell": cell.coordinate}
                if cell.comment:
                    comments += 1
                    line += f" [comment={cell.comment.text}]"
                    metadata["comment"] = cell.comment.text
                lines.append(line)
                units.append(Unit(locator, relative, "xlsx-cell", value, metadata=metadata))
    return "\n".join(lines), units, {"sheets": workbook.sheetnames, "comments": comments}


def _ticket_units(text: str, relative: str) -> list[Unit]:
    ticket = Path(relative).stem
    units = [Unit(f"{ticket}#header", relative, "ticket-header", "\n".join(text.splitlines()[:6]))]
    for line in text.splitlines():
        match = TICKET_COMMENT.match(line.strip())
        if not match:
            continue
        month_key = unicodedata.normalize("NFKD", match.group("month")).encode("ascii", "ignore").decode("ascii").lower()[:4]
        month = MONTHS[month_key]
        timestamp = f"2026-{month:02d}-{int(match.group('day')):02d}T{match.group('time').zfill(5)}"
        units.append(Unit(f"{ticket}#{timestamp}", relative, "ticket-comment", match.group("text"), metadata={"timestamp": timestamp}))
    return units


def _text_units(text: str, relative: str) -> list[Unit]:
    stem = Path(relative).stem
    if re.fullmatch(r"(?:ACC|DATA|INT|OPS|PERF|SEC)-\d+", stem):
        return _ticket_units(text, relative)
    units: list[Unit] = []
    is_meeting = re.match(r"M\d{2}_", stem) is not None
    is_teams = stem.startswith("Teams_")
    doc_id = stem.split("_")[0] if is_meeting else stem
    for number, line in enumerate(text.splitlines(), 1):
        stripped = line.strip()
        if not stripped:
            continue
        match = (MEETING_TIME if is_meeting else TEAMS_TIME).match(stripped) if (is_meeting or is_teams) else None
        if match:
            hhmm = match.group("time").zfill(5)
            units.append(Unit(f"{doc_id}@{hhmm}", relative, "meeting-line" if is_meeting else "teams-message", match.group("text"), metadata={"line": number, "timestamp": hhmm}))
        else:
            units.append(Unit(f"{stem}#L{number:04d}", relative, "text-line", stripped, metadata={"line": number}))
    return units


IMAGE_TRANSCRIPTIONS: dict[str, str] = {
    "OPS-601_runbook.png": (
        "Transcription vérifiée de la capture du runbook (version affichée : 25 septembre 2026).\n\n"
        "1. Vérifier la santé des services — OK.\n2. Activer le mode maintenance — OK.\n3. Déployer la version approuvée — OK.\n"
        "4. Procédure de retour arrière — TODO.\n5. Validation fonctionnelle post-déploiement — À compléter.\n"
    ),
}


def _extract_image(raw: bytes, relative: str, images_dir: Path, transcriptions_dir: Path) -> tuple[str, list[Unit], dict[str, Any]]:
    name = Path(relative).name
    (images_dir / name).write_bytes(raw)
    transcription = IMAGE_TRANSCRIPTIONS.get(
        name,
        "Transcription manuelle attendue. La capture est conservée intégralement; ne pas l'utiliser seule pour déduire un statut courant.\n",
    )
    transcription_path = transcriptions_dir / f"{name}.transcription.md"
    transcription_path.write_text(transcription, encoding="utf-8", newline="\n")
    units = [Unit(f"{name}#image", relative, "image", transcription)]
    if name == "OPS-601_runbook.png":
        for step, text in enumerate([
            "Vérifier la santé des services — OK", "Activer le mode maintenance — OK", "Déployer la version approuvée — OK",
            "Procédure de retour arrière — TODO", "Validation fonctionnelle post-déploiement — À compléter",
        ], 1):
            units.append(Unit(f"{name}#step{step}", relative, "image-step", text, metadata={"step": step}))
    return transcription, units, {"copied_image": f"images/{name}", "transcription": f"transcriptions/{name}.transcription.md", "transcription_status": "verified" if name in IMAGE_TRANSCRIPTIONS else "expected"}


def _render_evidence(relative: str, units: Iterable[Unit], evidence_dir: Path, source_asset: str | None = None) -> str:
    output_name = evidence_name(relative)
    blocks = []
    for unit in units:
        blocks.append(
            f'<section id="{html.escape(unit.anchor or anchor_for(unit.id), quote=True)}">'
            f'<h2><code>{html.escape(unit.id)}</code></h2><pre>{html.escape(unit.text)}</pre></section>'
        )
    asset = f'<p><a href="{html.escape(source_asset, quote=True)}">Ouvrir le fichier original</a></p>' if source_asset else ""
    document = f"""<!doctype html>
<html lang="fr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>Preuve — {html.escape(relative)}</title><style>
body{{font:16px/1.5 system-ui,sans-serif;max-width:72rem;margin:auto;padding:2rem}}pre{{white-space:pre-wrap}}
section{{border-top:1px solid #ccd;padding:1rem}}:target{{background:#fff3a3;outline:2px solid #e0b000;scroll-margin-top:2rem}}
</style></head><body><h1>{html.escape(relative)}</h1>{asset}{''.join(blocks)}</body></html>"""
    (evidence_dir / output_name).write_text(document, encoding="utf-8", newline="\n")
    return f"evidence/{output_name}"


def _read_manifest(archive: zipfile.ZipFile) -> tuple[dict[str, dict[str, str]], list[str]]:
    raw = archive.read(CORPUS_PREFIX + "MANIFEST.csv")
    text, _ = decode_text(raw)
    rows = list(csv.DictReader(io.StringIO(text)))
    if not rows or set(rows[0]) != {"fichier", "extension", "taille_octets"}:
        raise ValueError("MANIFEST.csv has an unexpected schema")
    manifest = {row["fichier"]: row for row in rows}
    errors = []
    if len(manifest) != len(rows):
        errors.append("MANIFEST.csv contains duplicate paths")
    return manifest, errors


def extract_corpus(zip_path: Path, output_root: Path, sources_json: Path | None = None) -> dict[str, Any]:
    """Extract all NOVA sources and return a machine-readable run report."""
    zip_path = Path(zip_path).resolve()
    output_root = Path(output_root).resolve()
    if not zip_path.is_file():
        raise FileNotFoundError(zip_path)
    for child in ("corpus", "text", "attachments", "images", "transcriptions", "evidence"):
        target = output_root / child
        if target.exists():
            shutil.rmtree(target)
        target.mkdir(parents=True, exist_ok=True)

    locators: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    warnings: list[str] = []
    with zipfile.ZipFile(zip_path) as archive:
        members = _safe_members(archive)
        manifest, manifest_errors = _read_manifest(archive)
        if manifest_errors:
            raise ValueError("; ".join(manifest_errors))
        actual_paths = {info.filename[len(CORPUS_PREFIX):] for info in members}
        expected_paths = set(manifest) | {"MANIFEST.csv"}
        if actual_paths != expected_paths:
            missing = sorted(expected_paths - actual_paths)
            extra = sorted(actual_paths - expected_paths)
            raise ValueError(f"Manifest mismatch; missing={missing}, extra={extra}")

        for info in members:
            relative = info.filename[len(CORPUS_PREFIX):]
            raw = archive.read(info)
            corpus_path = output_root / "corpus" / Path(*PurePosixPath(relative).parts)
            corpus_path.parent.mkdir(parents=True, exist_ok=True)
            corpus_path.write_bytes(raw)
            suffix = PurePosixPath(relative).suffix.lower()
            metadata: dict[str, Any] = {}
            encoding: str | None = None
            if suffix == ".eml":
                text, units, metadata = _extract_email(raw, relative, output_root / "attachments")
            elif suffix == ".pdf":
                text, units, metadata = _extract_pdf(raw, relative)
            elif suffix == ".xlsx":
                text, units, metadata = _extract_xlsx(raw, relative)
            elif suffix == ".png":
                text, units, metadata = _extract_image(raw, relative, output_root / "images", output_root / "transcriptions")
            else:
                text, encoding = decode_text(raw)
                units = _text_units(text, relative)
            if not units:
                raise ValueError(f"No addressable evidence units generated for {relative}")

            text_path = output_root / "text" / Path(*PurePosixPath(relative + ".txt").parts)
            text_path.parent.mkdir(parents=True, exist_ok=True)
            text_path.write_text(text, encoding="utf-8", newline="\n")
            evidence_file = _render_evidence(relative, units, output_root / "evidence", f"../corpus/{relative}")
            locators.extend(unit.as_dict(evidence_file) for unit in units)
            classification, reason, duplicate_group = _source_class(relative)
            manifest_row = manifest.get(relative)
            declared_size = int(manifest_row["taille_octets"]) if manifest_row else None
            if declared_size is not None and declared_size != len(raw):
                warnings.append(f"Manifest size differs for {relative}: declared={declared_size}, actual={len(raw)}")
            sources.append({
                "path": relative,
                "extension": suffix,
                "bytes": len(raw),
                "sha256": sha256(raw),
                "classification": classification,
                "classification_reason": reason,
                "information_date": _information_date(relative),
                "duplicate_group": duplicate_group,
                "extraction_status": "ok",
                "encoding": encoding,
                "text_path": f"text/{relative}.txt",
                "evidence_path": evidence_file,
                "locator_count": len(units),
                **metadata,
            })

    # Annotate standalone files that are byte-identical to saved attachments.
    attachment_hashes: dict[str, list[str]] = {}
    for source in sources:
        for attachment in source.get("attachments", []):
            attachment_hashes.setdefault(attachment["sha256"], []).append(f"{source['path']}::{attachment['name']}")
    for source in sources:
        matches = attachment_hashes.get(source["sha256"], [])
        if matches:
            source["attachment_copies"] = matches
            source["duplicate_group"] = source["duplicate_group"] or f"attachment-{source['sha256'][:12]}"

    # The archived 17 Sept email is intentionally byte-identical to E12.
    email_bodies = {item["id"]: item["text"] for item in locators if item["kind"] == "email-body"}
    if email_bodies.get("E12#body") != email_bodies.get("Courriel#body"):
        # Archive doc id derives from its filename; locate it by file instead.
        archive_body = next((item["text"] for item in locators if item["file"].endswith("Courriel_archive_17sept.eml") and item["kind"] == "email-body"), None)
        if archive_body != email_bodies.get("E12#body"):
            raise ValueError("Expected archive email duplicate does not match E12")

    locators.sort(key=lambda item: (item["file"], item["id"], item["anchor"]))
    sources.sort(key=lambda item: item["path"])
    locator_path = output_root / "locators.jsonl"
    locator_path.write_text("".join(json.dumps(item, ensure_ascii=False, sort_keys=True) + "\n" for item in locators), encoding="utf-8", newline="\n")
    report = {
        "schema_version": 1,
        "zip_sha256": sha256(zip_path.read_bytes()),
        "source_count": len(sources),
        "locator_count": len(locators),
        "class_counts": dict(sorted(__import__("collections").Counter(s["classification"] for s in sources).items())),
        "warnings": warnings,
        "sources": sources,
    }
    (output_root / "inventory.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n")
    if sources_json:
        sources_json = Path(sources_json)
        sources_json.parent.mkdir(parents=True, exist_ok=True)
        sources_json.write_text(json.dumps(sources, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n")
    return report
