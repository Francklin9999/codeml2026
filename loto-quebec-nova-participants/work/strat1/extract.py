#!/usr/bin/env python
"""Step A: corpus extraction shared by every NOVA strategy (strat1.md section 4.2).

    .venv/Scripts/python.exe -X utf8 loto-quebec-nova-participants/work/strat1/extract.py

Produces (all under work/_local/, git-ignored, fully regenerable):
    corpus/        the unzipped files (top folder 'Projet360_NOVA_ETUDIANTS/' stripped; never edit)
    text/<rel>.txt UTF-8 text for EVERY corpus file (rel keeps the original extension: foo.pdf -> foo.pdf.txt)
    attachments/   bytes of e-mail attachments, named <eml_stem>__<filename>
    images/        copy of the 8 PNG screenshots
    locators.jsonl one JSON record per addressable unit (see work/strat1/SCHEMA.md)
    files_index.json  one record per corpus file (rel, doc_id, folder class, sha256, text path)
    attachments_manifest.json  sha256 of each e-mail attachment vs the standalone corpus files
PNG screenshots have no machine text: their text/ file is generated from the hand-written
work/strat1/transcriptions/<png_stem>.md (placeholder if missing).
"""
import email
import email.policy
import email.utils
import hashlib
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parent))
import nova_common as C  # noqa: E402

MONTHS = {"janvier": 1, "janv": 1, "février": 2, "fevrier": 2, "févr": 2, "mars": 3, "avril": 4, "avr": 4,
          "mai": 5, "juin": 6, "juillet": 7, "juil": 7, "août": 8, "aout": 8, "septembre": 9, "sept": 9,
          "octobre": 10, "oct": 10, "novembre": 11, "nov": 11, "décembre": 12, "déc": 12}
FOLDER_CLASS = {"01_Courriels": "email", "02_Reunions": "meeting", "03_Tickets": "ticket",
                "04_Documents_projet": "project_doc", "05_Contrats_et_finances": "contract_finance",
                "06_Architecture_et_decisions": "architecture_decision", "07_Conversations_Teams": "teams",
                "08_Archives_et_documents_connexes": "archive_distractor"}

RECORDS: list = []
SEEN: dict = {}
STATS = {"decode_fallback": 0, "decode_errors": 0}


def emit(id_, file, kind, text, **extra):
    if id_ in SEEN:
        raise ValueError(f"duplicate locator id {id_} ({file} vs {SEEN[id_]})")
    SEEN[id_] = file
    RECORDS.append({"id": id_, "file": file, "kind": kind, "text": text.strip(), **extra})


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def decode(raw: bytes, rel: str) -> str:
    try:
        s = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        STATS["decode_fallback"] += 1
        try:
            s = raw.decode("cp1252")
        except UnicodeDecodeError:
            STATS["decode_errors"] += 1
            s = raw.decode("utf-8", errors="replace")
    return s.replace("\r\n", "\n").replace("\r", "\n")


def doc_id_for(rel: str) -> str:
    """Short document id used in locators: M04, E05, SEC-210, Teams_19sept_Securite; else the basename."""
    p = Path(rel)
    stem = p.stem
    m = re.match(r"^(M\d{2}|E\d{2})_", p.name)
    if m:
        return m.group(1)
    if rel.startswith("03_Tickets/") and p.suffix == ".txt" and re.fullmatch(r"[A-Z]+-\d{3}", stem):
        return stem
    if rel.startswith("07_Conversations_Teams/"):
        return stem
    return p.name


def doc_date(lines):
    """First 'D month YYYY' in the first lines -> (YYYY, MM, DD) or None."""
    for ln in lines[:4]:
        m = re.search(r"(\d{1,2})(?:er)?\s+([^\W\d_]+)\s+(\d{4})", ln)
        if m and m.group(2).lower() in MONTHS:
            return int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1))
    return None


def with_ordinals(counter: dict, key: str, base: str) -> str:
    """Duplicate-safe id: unique ids stay bare, repeated ones are resolved later with ~n (see finalize_dups)."""
    counter.setdefault(key, []).append(base)
    return base


def finalize_dups(units):
    """units: list of (base_id, record_kwargs). Repeated base ids become base~1, base~2 ... (all of them)."""
    count = {}
    for bid, _ in units:
        count[bid] = count.get(bid, 0) + 1
    seen = {}
    for bid, kw in units:
        if count[bid] > 1:
            seen[bid] = seen.get(bid, 0) + 1
            kw = dict(kw, ambiguous_base=bid)
            yield f"{bid}~{seen[bid]}", kw
        else:
            yield bid, kw


# ---------------------------------------------------------------- text-like files
def lines_units(rel, text, did):
    """#L<n> for every non-empty line (all text-like files)."""
    for n, ln in enumerate(text.split("\n"), 1):
        if ln.strip():
            emit(f"{did}#L{n}", rel, "text_line", ln, line=n)


def parse_meeting(rel, text, did):
    lines = text.split("\n")
    d = doc_date(lines)
    teams = rel.startswith("07_Conversations_Teams/")
    pat = re.compile(r"^(\d{2}:\d{2})\s+-\s+(.+?)\s*:\s*(.*)$") if teams else re.compile(r"^(\d{2}:\d{2})\s+(.*)$")
    units = []
    for n, ln in enumerate(lines, 1):
        m = pat.match(ln)
        if not m:
            continue
        hhmm = m.group(1)
        if teams:
            speaker, said = m.group(2), m.group(3)
        else:
            sm = re.match(r"^(.+?)\s*:\s*(.*)$", m.group(2))
            speaker, said = (sm.group(1), sm.group(2)) if sm else ("", m.group(2))
        ts = f"{d[0]:04d}-{d[1]:02d}-{d[2]:02d}T{hhmm}" if d else None
        units.append((f"{did}@{hhmm}", dict(file=rel, kind="teams_line" if teams else "meeting_line", text=ln,
                                             line=n, time=hhmm, speaker=speaker, said=said, ts=ts)))
    for id_, kw in finalize_dups(units):
        f = kw.pop("file"), kw.pop("kind"), kw.pop("text")
        emit(id_, f[0], f[1], f[2], **kw)


FIELD = re.compile(r"^(Titre|Créé|Demandeur|Priorité|Statut|Pièce jointe|Description initiale|Description|Contexte|"
                   r"Preuve|Analyse|Correction|Validation|Étapes|Attendu|Observé|Résolution) :\s*(.*)$")
COMMENT = re.compile(r"^(\d{1,2})\s+([^\W\d_]+)\.?\s+(?:(\d{1,2}:\d{2})\s+)?-\s+(.+?)\s*:\s*(.*)$")
DATED = re.compile(r"^(\d{1,2})\s+([^\W\d_]+)\s+-\s+(.*)$")


def slug(s):
    import unicodedata
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def parse_ticket(rel, text, did):
    lines = text.split("\n")
    year = 2026
    for ln in lines:
        m = re.match(r"^Créé : \d{1,2} [^\W\d_]+ (\d{4})", ln)
        if m:
            year = int(m.group(1))
    first_comment = next((i for i, ln in enumerate(lines) if ln.startswith("Commentaires :")), len(lines))
    emit(f"{did}#header", rel, "ticket_header", "\n".join(lines[:first_comment]), line=1)
    cur = None  # current field accumulator
    fields, comment_units, in_comments = [], [], False

    def flush():
        nonlocal cur
        if cur:
            fields.append(cur)
            cur = None

    for n, ln in enumerate(lines, 1):
        if ln.startswith("Commentaires :"):
            flush()
            in_comments = True
            continue
        fm = FIELD.match(ln)
        if fm:
            flush()
            in_comments = False if fm.group(1) in ("Résolution",) else in_comments
            cur = {"name": fm.group(1), "lines": [fm.group(2)] if fm.group(2) else [], "line": n}
            continue
        cm = COMMENT.match(ln) if in_comments else None
        if cm:
            flush()
            mon = MONTHS.get(cm.group(2).lower())
            if mon is None:
                raise ValueError(f"{rel}:{n} unknown month {cm.group(2)}")
            ts = f"{year:04d}-{mon:02d}-{int(cm.group(1)):02d}" + (f"T{int(cm.group(3).split(':')[0]):02d}:{cm.group(3).split(':')[1]}" if cm.group(3) else "")
            comment_units.append((f"{did}#{ts}", dict(file=rel, kind="ticket_comment", text=ln, line=n, ts=ts,
                                                      speaker=cm.group(4), said=cm.group(5))))
            continue
        dm = DATED.match(ln)
        if dm and dm.group(2).lower() in MONTHS and not in_comments:
            ts = f"{year:04d}-{MONTHS[dm.group(2).lower()]:02d}-{int(dm.group(1)):02d}"
            comment_units.append((f"{did}#{ts}", dict(file=rel, kind="ticket_dated_line", text=ln, line=n, ts=ts)))
            if cur:
                cur["lines"].append(ln)
            continue
        if cur is not None:
            if ln.strip():
                cur["lines"].append(ln)
            else:
                flush()
    flush()
    for f in fields:
        emit(f"{did}#{slug(f['name'])}", rel, "ticket_field", "\n".join(f["lines"]) or "(vide)", line=f["line"],
             field=f["name"])
    for id_, kw in finalize_dups(comment_units):
        f = kw.pop("file"), kw.pop("kind"), kw.pop("text")
        emit(id_, f[0], f[1], f[2], **kw)


def process_text_file(rel, raw):
    text = decode(raw, rel)
    did = doc_id_for(rel)
    name = Path(rel).name
    if rel.startswith(("02_Reunions/", "07_Conversations_Teams/")):
        parse_meeting(rel, text, did)
    elif rel.startswith("03_Tickets/") and re.fullmatch(r"[A-Z]+-\d{3}\.txt", name):
        parse_ticket(rel, text, did)
    elif name == "README.txt":
        for n, ln in enumerate(text.split("\n"), 1):
            m = re.match(r"^(Q\d{2})\.", ln)
            if m:
                emit(f"README.txt#{m.group(1)}", rel, "readme_question", ln, line=n)
    lines_units(rel, text, did)
    return text


# ---------------------------------------------------------------- e-mail
def pdf_text_pages(src):
    doc = pymupdf.open(stream=src, filetype="pdf") if isinstance(src, (bytes, bytearray)) else pymupdf.open(src)
    pages = []
    for p in doc:
        pages.append((p.get_text("text").strip(), pdf_rows(p)))
    return pages


def pdf_rows(page):
    """Visual rows rebuilt from word coordinates: cells of a row joined by ' | '."""
    words = sorted(page.get_text("words"), key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    rows, cur, cy = [], [], None
    for w in words:
        y = (w[1] + w[3]) / 2
        if cy is None or abs(y - cy) <= 3:
            cur.append(w)
            cy = y if cy is None else cy
        else:
            rows.append(cur)
            cur, cy = [w], y
    if cur:
        rows.append(cur)
    out = []
    for r in rows:
        r.sort(key=lambda w: w[0])
        s, last = r[0][4], r[0]
        for w in r[1:]:
            s += (" | " if w[0] - last[2] > 12 else " ") + w[4]
            last = w
        out.append(s)
    return out


def process_eml(rel, raw, did):
    msg = email.message_from_bytes(raw, policy=email.policy.default)
    hdr = {k: str(msg[k]) for k in ("From", "To", "Date", "Subject", "Message-ID") if msg[k]}
    body_part = msg.get_body(preferencelist=("plain",))
    body = body_part.get_content().replace("\r\n", "\n").strip() if body_part else ""
    atts = []
    stem = Path(rel).stem
    C.ATTACH.mkdir(parents=True, exist_ok=True)
    for att in msg.iter_attachments():
        fn = att.get_filename() or "attachment.bin"
        data = att.get_content()
        data = data if isinstance(data, (bytes, bytearray)) else str(data).encode("utf-8")
        saved = C.ATTACH / f"{stem}__{fn}"
        saved.write_bytes(data)
        info = {"eml": rel, "filename": fn, "saved_as": saved.name, "size": len(data), "sha256": sha256(data),
                "content_type": att.get_content_type()}
        if fn.lower().endswith(".pdf"):
            pages = pdf_text_pages(bytes(data))
            info["text"] = "\n".join(f"=== page {i} ===\n{t}" for i, (t, _) in enumerate(pages, 1))
        atts.append(info)
    try:
        iso = email.utils.parsedate_to_datetime(hdr.get("Date", "")).isoformat()
    except Exception:
        iso = None
    htxt = "\n".join(f"{k}: {v}" for k, v in hdr.items())
    parts = [htxt, "", "--- body ---", body]
    if atts:
        parts += ["", "--- attachments ---"] + [f"- {a['filename']} ({a['content_type']}, {a['size']} bytes, sha256 {a['sha256'][:12]})" for a in atts]
        for a in atts:
            if a.get("text"):
                parts += ["", f"--- attachment text: {a['filename']} ---", a["text"]]
    emit(f"{did}#headers", rel, "email_headers", htxt, date_iso=iso, sender=hdr.get("From"), subject=hdr.get("Subject"))
    emit(f"{did}#body", rel, "email_body", body, date_iso=iso)
    for a in atts:
        emit(f"{did}#att:{a['filename']}", rel, "email_attachment", a.get("text", ""), sha256=a["sha256"], saved_as=a["saved_as"])
    return "\n".join(parts) + "\n", atts


# ---------------------------------------------------------------- pdf / xlsx / png
def process_pdf(rel, path):
    name = Path(rel).name
    pages = pdf_text_pages(path)
    out = []
    for i, (t, rows) in enumerate(pages, 1):
        out.append(f"=== page {i} ===\n{t}")
        emit(f"{name}#p{i}", rel, "pdf_page", t, page=i)
        for j, r in enumerate(rows, 1):
            emit(f"{name}#p{i}/r{j}", rel, "pdf_row", r, page=i, row=j)
    return "\n".join(out) + "\n"


def process_xlsx(rel, path):
    name = Path(rel).name
    wb = openpyxl.load_workbook(path, data_only=True)
    multi = len(wb.sheetnames) > 1
    out, ncomments = [], 0
    for ws in wb.worksheets:
        pre = f"{name}!{ws.title}!" if multi else f"{name}!"
        out.append(f"# sheet: {ws.title} ({ws.dimensions})")
        for row in ws.iter_rows():
            cells = []
            for c in row:
                if c.comment:
                    ncomments += 1
                    emit(f"{pre}{c.coordinate}#comment", rel, "xlsx_comment", c.comment.text, cell=c.coordinate)
                if c.value is None or str(c.value).strip() == "":
                    continue
                cells.append(f"{c.coordinate}={c.value}")
                emit(f"{pre}{c.coordinate}", rel, "xlsx_cell", str(c.value), cell=c.coordinate, sheet=ws.title,
                     row=c.row, col=c.column_letter)
            if cells:
                line = " | ".join(cells)
                out.append(line)
                emit(f"{pre}row{row[0].row}", rel, "xlsx_row", line, sheet=ws.title, row=row[0].row)
    out.append(f"# cell comments found: {ncomments}")
    return "\n".join(out) + "\n"


def process_png(rel, path):
    shutil.copy2(path, C.IMAGES / path.name)
    stem = Path(rel).stem
    tr = C.TRANSCRIPTIONS / f"{stem}.md"
    name = Path(rel).name
    if tr.exists():
        body = tr.read_text(encoding="utf-8").replace("\r\n", "\n")
        emit(f"{name}#img", rel, "image", body, transcription=str(tr.relative_to(C.ROOT).as_posix()))
        for m in re.finditer(r"^\s*[-*]\s*`#([\w.-]+)`\s*[:|]\s*(.*)$", body, re.M):
            emit(f"{name}#{m.group(1)}", rel, "image_unit", m.group(2))
        return body
    return f"[IMAGE {name}: no transcription file yet ({tr.name})]\n"


# ---------------------------------------------------------------- main
def main():
    for d in (C.CORPUS, C.TEXT, C.ATTACH, C.IMAGES):
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True)
    # 1. unzip (strip the top folder; refuse paths escaping the corpus dir)
    with zipfile.ZipFile(C.ZIP_PATH) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            rel = info.filename[len(C.TOP_DIR):] if info.filename.startswith(C.TOP_DIR) else info.filename
            dest = (C.CORPUS / rel).resolve()
            if C.CORPUS.resolve() not in dest.parents:
                raise ValueError(f"unsafe zip path {info.filename}")
            dest.parent.mkdir(parents=True, exist_ok=True)
            dest.write_bytes(z.read(info))
    files = sorted(p for p in C.CORPUS.rglob("*") if p.is_file())
    index, n_text, n_img = [], 0, 0
    for p in files:
        rel = p.relative_to(C.CORPUS).as_posix()
        raw = p.read_bytes()
        ext = p.suffix.lower()
        did = doc_id_for(rel)
        if ext == ".eml":
            text, _ = process_eml(rel, raw, did)
        elif ext == ".pdf":
            text = process_pdf(rel, p)
        elif ext == ".xlsx":
            text = process_xlsx(rel, p)
        elif ext == ".png":
            text = process_png(rel, p)
            n_img += 1
        else:
            text = process_text_file(rel, raw)
        out = C.text_path(rel)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text, encoding="utf-8", newline="\n")
        n_text += 1
        index.append({"file": rel, "doc_id": did, "ext": ext, "folder": rel.split("/")[0] if "/" in rel else "(root)",
                      "class": FOLDER_CLASS.get(rel.split("/")[0], "root"), "size": len(raw), "sha256": sha256(raw),
                      "text": C.text_path(rel).relative_to(C.ROOT).as_posix()})
    # 2. outputs
    with open(C.LOCATORS, "w", encoding="utf-8", newline="\n") as fh:
        for r in RECORDS:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    (C.LOCAL / "files_index.json").write_text(json.dumps(index, ensure_ascii=False, indent=1), encoding="utf-8")
    # attachment duplicates: compare sha256 to every corpus file
    by_hash = {}
    for e in index:
        by_hash.setdefault(e["sha256"], []).append(e["file"])
    man = []
    for f in sorted(C.ATTACH.iterdir()):
        h = sha256(f.read_bytes())
        man.append({"attachment": f.name, "sha256": h, "size": f.stat().st_size, "identical_to_corpus_files": by_hash.get(h, [])})
    (C.LOCAL / "attachments_manifest.json").write_text(json.dumps(man, ensure_ascii=False, indent=1), encoding="utf-8")
    kinds = {}
    for r in RECORDS:
        kinds[r["kind"]] = kinds.get(r["kind"], 0) + 1
    print(f"corpus files: {len(files)} | text files: {n_text} | images: {n_img} | attachments: {len(man)}")
    print(f"locators: {len(RECORDS)} | decode fallbacks: {STATS['decode_fallback']} | decode errors: {STATS['decode_errors']}")
    print("locators by kind:", json.dumps(kinds, ensure_ascii=False, sort_keys=True))
    print("attachments identical to a standalone corpus file:",
          sum(1 for m in man if m["identical_to_corpus_files"]), "/", len(man))
    return 0 if STATS["decode_errors"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
